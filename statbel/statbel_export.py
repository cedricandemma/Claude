#!/usr/bin/env python3
"""Génère les fichiers CSV d'upload Statbel (enquête Tourisme et hôtellerie, nuitées)
à partir des exports de réservations du site (MotoPress Hotel Booking, .csv)
et de Booking.com (.xls).

Usage :
    python3 statbel_export.py <dossier_exports> [--mois 2026-04] [--out <dossier_sortie>]

Un fichier statbel_AAAA-MM.csv est produit par mois, au format du modèle
Tourism_template_XBRL.xlsx enregistré en CSV (séparateur point-virgule, sans titres) :
    pays ; but du séjour ; jour de départ (aaaa-mm-jj) ; nuits ; personnes ; unités ; région BE
Statbel rattache un séjour au mois de sa date de départ (check-out).
Le websurvey exige 7 colonnes (le modèle Excel n'en montre que 6, il est antérieur à la
colonne région). La région n'est remplie que pour les résidents belges : d'après le code
postal quand il est connu, sinon selon la répartition 80 % Flandre / 20 % Wallonie.
"""
import argparse
import csv
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import xlrd

# Codes du websurvey Statbel.
# But du séjour : 1 loisirs, vacances ; 2 conférence, congrès, séminaire ;
# 3 autres raisons professionnelles. Booking.com ne distingue que "Affaires" -> 3.
BUT_LOISIRS = "1"
BUT_PROFESSIONNEL = "3"
# Région de résidence en Belgique.
REGION_WALLONIE = "1"
REGION_BRUXELLES = "2"
REGION_FLANDRE = "3"
REGION_GERMANOPHONE = "4"
# Clientèle belge sans code postal connu : 4 sur 5 en Flandre, 1 sur 5 en Wallonie.
PART_WALLONIE_SANS_CP = 5

# Onglet CountryCode du modèle Statbel ; tout autre code est déclaré XX (indéterminé).
PAYS_STATBEL = set("""
AD AE AF AG AI AL AM AO AQ AR AS AT AU AW AX AZ BA BB BD BE BF BG BH BI BJ BL BM BN BO BQ
BR BS BT BV BW BY BZ CA CC CD CF CG CH CI CK CL CM CN CO CR CU CV CW CX CY CZ DE DJ DK DM
DO DZ EC EE EG EH ER ES ET FI FJ FK FM FO FR GA GB GD GE GF GG GH GI GL GM GN GP GQ GR GS
GT GU GW GY HK HM HN HR HT HU ID IE IL IM IN IO IQ IR IS IT JE JM JO JP KE KG KH KI KM KN
KP KR KW KY KZ LA LB LC LI LK LR LS LT LU LV LY MA MC MD ME MF MG MH MK ML MM MN MO MP MQ
MR MS MT MU MV MW MX MY MZ NA NC NE NF NG NI NL NO NP NR NU NZ OM PA PE PF PG PH PK PL PM
PN PR PS PT PW PY QA RE RO RS RU RW SA SB SC SD SE SG SH SI SJ SK SL SM SN SO SR SS ST SV
SX SY SZ TC TD TF TG TH TJ TK TL TM TN TO TR TT TV TW TZ UA UG UM US UY UZ VA VC VE VG VI
VN VU WF WS XX YE YT ZA ZM ZW
""".split())
ALIAS_PAYS = {"UK": "GB", "EL": "GR"}


def code_pays(brut):
    c = str(brut).strip().upper()
    c = ALIAS_PAYS.get(c, c)
    return c if c in PAYS_STATBEL else "XX"

# Hypothèse pour le site web, qui n'exporte pas le nombre d'occupants :
# chambres doubles, donc 2 personnes par chambre.
PERSONNES_PAR_CHAMBRE_SITE = 2


def region_depuis_code_postal(cp):
    m = re.search(r"\b(\d{4})\b", cp or "")
    if not m:
        return ""
    n = int(m.group(1))
    if 1000 <= n <= 1299:
        return REGION_BRUXELLES
    if 4700 <= n <= 4799:
        return REGION_GERMANOPHONE
    if 1300 <= n <= 1499 or 4000 <= n <= 7999:
        return REGION_WALLONIE
    if 1500 <= n <= 3999 or 8000 <= n <= 9999:
        return REGION_FLANDRE
    return ""


def lire_site(path):
    """Export MotoPress : une ligne par chambre, même ID si plusieurs chambres."""
    sejours = {}
    with open(path, encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["Status"].strip().lower() != "confirmed":
                continue
            arrivee = datetime.strptime(row["Check-in"], "%d/%m/%Y").date()
            depart = datetime.strptime(row["Check-out"], "%d/%m/%Y").date()
            pays = code_pays(row["Country"])
            s = sejours.setdefault(row["ID"], {
                "source": "site", "ref": row["ID"], "nom": f'{row["First Name"]} {row["Last Name"]}',
                "pays": pays, "but": BUT_LOISIRS, "depart": depart,
                "nuits": (depart - arrivee).days, "unites": 0, "personnes": 0,
                "region": region_depuis_code_postal(row["Postcode"]) if pays == "BE" else "",
            })
            s["unites"] += 1
            s["personnes"] += PERSONNES_PAR_CHAMBRE_SITE
    return list(sejours.values())


def ouvrir_xls(path):
    data = Path(path).read_bytes()
    try:
        return xlrd.open_workbook(file_contents=data)
    except AssertionError:
        # Certains exports Booking.com ont un conteneur OLE mal formé :
        # le flux BIFF commence juste après l'en-tête de 512 octets.
        return xlrd.open_workbook(file_contents=data[512:], ignore_workbook_corruption=True)


def lire_booking(path):
    sh = ouvrir_xls(path).sheet_by_index(0)
    entetes = [h.replace("&nbsp;", "").strip() for h in sh.row_values(0)]
    sejours = []
    for r in range(1, sh.nrows):
        v = dict(zip(entetes, sh.row_values(r)))
        if str(v["Statut"]).strip().lower() != "ok":
            continue
        arrivee = datetime.strptime(v["Arrivée"], "%Y-%m-%d").date()
        depart = datetime.strptime(v["Départ"], "%Y-%m-%d").date()
        pays = code_pays(v["Booker country"])
        motif = str(v["Motif du voyage"]).strip().lower()
        sejours.append({
            "source": "booking", "ref": str(v["Numéro de réservation"]).split(".")[0],
            "nom": v["Nom du client"], "pays": pays,
            "but": BUT_PROFESSIONNEL if motif == "affaires" else BUT_LOISIRS,
            "depart": depart, "nuits": (depart - arrivee).days,
            "unites": int(v["Hébergements"] or 1), "personnes": int(v["Personnes"] or 2),
            "region": region_depuis_code_postal(str(v.get("Adresse", ""))) if pays == "BE" else "",
        })
    return sejours


LIBELLES_BUT = {"1": "(1) loisirs, vacances", "2": "(2) conférence, congrès, séminaire",
                "3": "(3) autres raisons professionnelles"}
LIBELLES_REGION = {"1": "(1) Wallonie", "2": "(2) Bruxelles", "3": "(3) Flandre",
                   "4": "(4) Communauté germanophone"}


def ecrire_classeur_saisie(chemin, mois_lignes):
    """Classeur lisible pour la saisie manuelle : un onglet par mois, libellés du websurvey."""
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for mois, lignes in mois_lignes:
        ws = wb.create_sheet(mois)
        ws.append(["Pays de résidence", "But du séjour", "Jour de départ", "Nombre de nuits",
                   "Nombre de personnes", "Nombre d'unités", "Région de résidence en Belgique",
                   "Source", "Réservation", "Client"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for s in lignes:
            ws.append([s["pays"], LIBELLES_BUT[s["but"]], s["depart"].strftime("%d/%m/%Y"),
                       s["nuits"], s["personnes"], s["unites"],
                       LIBELLES_REGION.get(s["region"], "") +
                       (" (estimée)" if s.get("region_estimee") else ""),
                       s["source"], s["ref"], s["nom"]])
        for col, largeur in zip("ABCDEFGHIJ", (10, 24, 14, 10, 12, 10, 30, 9, 13, 28)):
            ws.column_dimensions[col].width = largeur
        ws.freeze_panes = "A2"
    wb.save(chemin)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("dossier")
    ap.add_argument("--mois", help="AAAA-MM ; par défaut tous les mois présents")
    ap.add_argument("--out", default=".")
    ap.add_argument("--format-date", default="%Y-%m-%d",
                    help="format de la date de départ, ex. %%d/%%m/%%Y (défaut %%Y-%%m-%%d)")
    args = ap.parse_args()

    sejours = []
    for p in sorted(Path(args.dossier).iterdir()):
        if p.suffix.lower() == ".csv" and not p.name.startswith("statbel_"):
            sejours += lire_site(p)
        elif p.suffix.lower() == ".xls":
            sejours += lire_booking(p)

    par_mois = defaultdict(list)
    for s in sejours:
        par_mois[s["depart"].strftime("%Y-%m")].append(s)

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    pour_saisie = []
    for mois in sorted(par_mois):
        if args.mois and mois != args.mois:
            continue
        lignes = sorted(par_mois[mois], key=lambda s: (s["depart"], s["ref"]))
        for s in lignes:
            if s["pays"] == "BE" and not s["region"]:
                # Répartition stable d'un mois à l'autre : 1 réservation sur 5 en Wallonie,
                # selon le dernier chiffre utile du numéro de réservation.
                s["region_estimee"] = True
                s["region"] = REGION_WALLONIE if int(s["ref"]) % PART_WALLONIE_SANS_CP == 0 \
                    else REGION_FLANDRE
        with open(out / f"statbel_{mois}.csv", "w", encoding="utf-8", newline="") as f:
            w = csv.writer(f, delimiter=";")
            for s in lignes:
                ligne = [s["pays"], s["but"], s["depart"].strftime(args.format_date), s["nuits"],
                         s["personnes"], s["unites"]]
                ligne.append(s["region"])
                w.writerow(ligne)
        pour_saisie.append((mois, lignes))
        nuitees = sum(s["nuits"] * s["personnes"] for s in lignes)
        print(f"\n{mois} : {len(lignes)} séjours, {nuitees} nuitées -> statbel_{mois}.csv")
        for s in lignes:
            alerte = "  <- région estimée" if s.get("region_estimee") else ""
            print(f'  départ {s["depart"]} {s["source"]:7} {s["ref"]:>11} {s["pays"]} '
                  f'{s["nuits"]}n {s["personnes"]}p {s["unites"]}ch {s["nom"]}{alerte}')

    if pour_saisie:
        ecrire_classeur_saisie(out / "statbel_saisie.xlsx", pour_saisie)
        print("\nClasseur de saisie manuelle -> statbel_saisie.xlsx")


if __name__ == "__main__":
    sys.exit(main())
